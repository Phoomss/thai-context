#!/usr/bin/env python3
"""
transform_dict_xlsx_to_json.py
Transform Dictionary XLSX files (2542, 2554, 2569) into clean, structured JSON files.
"""

import os
import re
import json
import openpyxl
from typing import Dict, Any, List, Optional

DICT_DIR = "data/processed/dict"

def clean_val(val: Any) -> Optional[str]:
    if val is None:
        return None
    s = str(val).strip()
    if s.lower() in ("null", "none", ""):
        return None
    return s

def extract_2542_components(dt_raw: Optional[str]) -> Dict[str, Any]:
    """
    In 2542, definition text (D_T) often starts with:
    - [อ่านว่า ...] or [โบ อ่านว่า ...]
    - (สํา) for idiom
    - POS: น., ก., ว., สัน., บ., ส., นิ., etc.
    """
    if not dt_raw:
        return {"pronunciation": None, "pos": None, "definition": ""}

    text = dt_raw.strip()
    pronunciation = None
    pos = None

    # Check for [อ่านว่า ...] or [โบ อ่านว่า ...]
    read_match = re.search(r"\[(?:[^\]]*\s)?อ่านว่า\s+([^\]]+)\]", text)
    if read_match:
        pronunciation = read_match.group(1).strip()
        text = text[:read_match.start()] + text[read_match.end():]
        text = text.strip()

    # Check for POS at the start of text (including vowels like สัน., ก.ว., อสํ., etc.)
    pos_match = re.match(r"^((?:\([^\)]+\)\s*)?(?:[ก-ฮ\u0e30-\u0e4e]+\.)+)\s*(.*)$", text)
    if pos_match:
        pos = pos_match.group(1).strip()
        text = pos_match.group(2).strip()

    return {
        "pronunciation": pronunciation,
        "pos": pos,
        "definition": text
    }

def clean_headword_and_sense(hw: str) -> tuple:
    """Split headword and sense indicator like 'ก็ ๑' -> ('ก็', 1)"""
    thai_digits = {"๑": 1, "๒": 2, "๓": 3, "๔": 4, "๕": 5, "๖": 6, "๗": 7, "๘": 8, "๙": 9, "๐": 0}
    m = re.search(r"\s+([๑-๙\d]+)$", hw)
    if m:
        digit_str = m.group(1)
        sense = 0
        for ch in digit_str:
            if ch in thai_digits:
                sense = sense * 10 + thai_digits[ch]
            elif ch.isdigit():
                sense = sense * 10 + int(ch)
        clean_hw = hw[:m.start()].strip()
        return clean_hw, sense if sense > 0 else 1
    return hw.strip(), 1

def transform_2569() -> List[Dict[str, Any]]:
    path = os.path.join(DICT_DIR, "DICT_2569 (ก_Incomplete).xlsx")
    print(f"🔄 Reading {path}...")
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    headers = rows[0]
    entries = []
    for idx, r in enumerate(rows[1:], start=1):
        hw = clean_val(r[0])
        if not hw:
            continue
        sense_raw = clean_val(r[1])
        pron = clean_val(r[2])
        pos = clean_val(r[3])
        definition = clean_val(r[4]) or ""

        clean_hw, inferred_sense = clean_headword_and_sense(hw)
        sense_num = int(sense_raw) if sense_raw and str(sense_raw).isdigit() else inferred_sense

        entries.append({
            "index": idx,
            "headword": clean_hw,
            "raw_headword": hw,
            "sense_order": sense_num,
            "pronunciation": pron,
            "pos": pos,
            "definition": definition,
            "edition": "2569",
            "edition_title": "พจนานุกรม ฉบับราชบัณฑิตยสถาน พ.ศ. ๒๕๖๙ (ฉบับดิจิทัล)"
        })

    out_path = os.path.join(DICT_DIR, "dict_2569.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(entries, f, ensure_ascii=False, indent=2)

    print(f"✅ Generated {out_path} ({len(entries):,} entries)")
    return entries

def transform_2554() -> List[Dict[str, Any]]:
    path = os.path.join(DICT_DIR, "DICT_2554 (ก_ซ).xlsx")
    print(f"🔄 Reading {path}...")
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    entries = []
    for idx, r in enumerate(rows[1:], start=1):
        hw = clean_val(r[1])
        if not hw:
            continue

        sense_raw = clean_val(r[2])
        pron = clean_val(r[3])
        subject_field = clean_val(r[4])
        pos = clean_val(r[5])
        definition = clean_val(r[6]) or ""
        book = clean_val(r[7])
        is_subword = clean_val(r[8]) == "1"
        usage = clean_val(r[9])
        etymology = clean_val(r[11])
        register = clean_val(r[12])
        sense_qualifier = clean_val(r[13])

        clean_hw, inferred_sense = clean_headword_and_sense(hw)
        sense_num = int(sense_raw) if sense_raw and str(sense_raw).isdigit() else inferred_sense

        entries.append({
            "index": idx,
            "headword": clean_hw,
            "raw_headword": hw,
            "sense_order": sense_num,
            "pronunciation": pron,
            "pos": pos,
            "definition": definition,
            "subject_field": subject_field,
            "is_subword": is_subword,
            "usage": usage,
            "etymology": etymology,
            "register": register,
            "sense_qualifier": sense_qualifier,
            "edition": "2554",
            "edition_title": "พจนานุกรม ฉบับราชบัณฑิตยสถาน พ.ศ. ๒๕๕๔"
        })

    out_path = os.path.join(DICT_DIR, "dict_2554.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(entries, f, ensure_ascii=False, indent=2)

    print(f"✅ Generated {out_path} ({len(entries):,} entries)")
    return entries

def transform_2542() -> List[Dict[str, Any]]:
    path = os.path.join(DICT_DIR, "DICT_2542 (ก_ฮ).xlsx")
    print(f"🔄 Reading {path}...")
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    entries = []
    for idx, r in enumerate(rows[1:], start=1):
        kw = clean_val(r[1])
        mw = clean_val(r[2])
        dt = clean_val(r[3])

        headword_source = mw or kw
        if not headword_source:
            continue

        clean_hw, inferred_sense = clean_headword_and_sense(headword_source)
        dt_parts = extract_2542_components(dt)

        entries.append({
            "index": idx,
            "headword": clean_hw,
            "raw_headword": headword_source,
            "keyword_variants": kw,
            "sense_order": inferred_sense,
            "pronunciation": dt_parts["pronunciation"],
            "pos": dt_parts["pos"],
            "definition": dt_parts["definition"],
            "edition": "2542",
            "edition_title": "พจนานุกรม ฉบับราชบัณฑิตยสถาน พ.ศ. ๒๕๔๒"
        })

    out_path = os.path.join(DICT_DIR, "dict_2542.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(entries, f, ensure_ascii=False, indent=2)

    print(f"✅ Generated {out_path} ({len(entries):,} entries)")
    return entries

def generate_evolution_and_combined(e2569: List[Dict], e2554: List[Dict], e2542: List[Dict]):
    # 1. Combined flat list
    all_entries = e2542 + e2554 + e2569
    combined_path = os.path.join(DICT_DIR, "dict_all_editions.json")
    print(f"🔄 Writing combined master JSON ({len(all_entries):,} entries)...")
    with open(combined_path, "w", encoding="utf-8") as f:
        json.dump(all_entries, f, ensure_ascii=False, indent=2)
    print(f"✅ Generated {combined_path}")

    # 2. Grouped evolution dictionary for multi-edition comparison
    merged: Dict[str, Dict[str, Any]] = {}
    for item in all_entries:
        hw = item["headword"]
        if hw not in merged:
            merged[hw] = {"headword": hw, "editions": {}}
        ed = item["edition"]
        if ed not in merged[hw]["editions"]:
            merged[hw]["editions"][ed] = []
        merged[hw]["editions"][ed].append({
            "sense_order": item.get("sense_order", 1),
            "pronunciation": item.get("pronunciation"),
            "pos": item.get("pos"),
            "definition": item.get("definition")
        })

    # Filter evolution words that appear in at least 2 editions
    evolution_words = {hw: data for hw, data in merged.items() if len(data["editions"]) >= 2}
    evolution_path = os.path.join(DICT_DIR, "dict_evolution_comparison.json")
    print(f"🔄 Writing evolution comparison JSON ({len(evolution_words):,} words)...")
    with open(evolution_path, "w", encoding="utf-8") as f:
        json.dump(evolution_words, f, ensure_ascii=False, indent=2)
    print(f"✅ Generated {evolution_path}")

def main():
    print("==================================================")
    print("🚀 Transforming Royal Society Dictionary XLSX to JSON")
    print("==================================================")

    e2569 = transform_2569()
    e2554 = transform_2554()
    e2542 = transform_2542()

    generate_evolution_and_combined(e2569, e2554, e2542)

    total_count = len(e2569) + len(e2554) + len(e2542)
    print("==================================================")
    print(f"🎉 All transformations complete!")
    print(f"- 2569 Edition: {len(e2569):,} entries")
    print(f"- 2554 Edition: {len(e2554):,} entries")
    print(f"- 2542 Edition: {len(e2542):,} entries")
    print(f"📊 Total JSON entries generated: {total_count:,}")
    print("==================================================")

if __name__ == "__main__":
    main()
